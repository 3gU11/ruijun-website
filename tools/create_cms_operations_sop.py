from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "瑞钧官网CMS运营标准化流程说明书_v1.0.docx"
ASSET_DIR = ROOT / "docs" / "operations" / "assets"
FLOW_DIAGRAM = ASSET_DIR / "cms-content-change-flow.png"
WORKBENCH_HOME = ASSET_DIR / "cms-editor-workbench-home.png"
WORKBENCH_TOOLS = ASSET_DIR / "cms-editor-workbench-tools.png"
PAGE_FORM = ASSET_DIR / "cms-editor-page-form.png"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
MUTED = "5B6776"
RED = "9B1C1C"
GOLD = "7A5A00"


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=TABLE_INDENT_DXA):
    if sum(widths) != PAGE_WIDTH_DXA:
        raise ValueError("table widths must sum to 9360 DXA")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color="C7D0DA", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def add_text(cell, text, size=10, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    mark_header_row(table.rows[0])
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        add_text(cell, header, size=10, bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            add_text(cell, value)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, label, text, color=INK):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CALLOUT)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "4")
        border.set(qn("w:color"), "D5DDE7")
        borders.append(border)
    p_pr.append(borders)
    r = p.add_run(f"{label}  ")
    set_run_font(r, size=10.5, color=color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color="27313D")


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        prefix = p.add_run(bold_prefix)
        set_run_font(prefix, size=11, color="000000", bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11, color="000000")
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color="000000")
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string("2B3A4B")
    return p


def add_figure(doc, path, caption, alt_text, width=6.25):
    if not path.exists():
        raise FileNotFoundError(f"Missing document figure: {path}")
    inline_shape = doc.add_picture(str(path), width=Inches(width))
    inline_shape._inline.docPr.set("descr", alt_text)
    inline_shape._inline.docPr.set("title", caption)
    image_paragraph = doc.paragraphs[-1]
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(4)
    image_paragraph.paragraph_format.space_after = Pt(3)
    caption_paragraph = doc.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(0)
    caption_paragraph.paragraph_format.space_after = Pt(9)
    run = caption_paragraph.add_run(caption)
    set_run_font(run, size=9.5, color=MUTED, italic=True)


def center_text(draw, box, text, font, fill, spacing=4):
    left, top, right, bottom = box
    draw.multiline_text(
        ((left + right) / 2, (top + bottom) / 2),
        text,
        font=font,
        fill=fill,
        anchor="mm",
        align="center",
        spacing=spacing,
    )


def create_flow_diagram():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    canvas = Image.new("RGB", (1800, 820), "#FFFFFF")
    draw = ImageDraw.Draw(canvas)
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    bold_font_path = Path("C:/Windows/Fonts/msyhbd.ttc")
    title_font = ImageFont.truetype(str(bold_font_path), 42)
    label_font = ImageFont.truetype(str(bold_font_path), 29)
    detail_font = ImageFont.truetype(str(font_path), 21)
    stage_font = ImageFont.truetype(str(bold_font_path), 18)
    note_font = ImageFont.truetype(str(font_path), 22)

    draw.rounded_rectangle((28, 26, 1772, 794), radius=28, fill="#F6F9FC", outline="#D9E3ED", width=3)
    center_text(draw, (80, 58, 1720, 114), "官网内容变更闭环", title_font, "#0B2545")
    center_text(draw, (80, 118, 1720, 150), "编辑人员维护草稿；审核、发布和恢复由相应角色在专用工作台执行", detail_font, "#5B6776")

    boxes = [
        ("STEP 01", "准备依据", "确认目标、来源和授权"),
        ("STEP 02", "编辑草稿", "填写内容与关联关系"),
        ("STEP 03", "双重预览", "表单预览 + 官网预览"),
        ("STEP 04", "提交审核", "写明变更与上线计划"),
        ("STEP 05", "审核与就绪", "处理意见并清除阻断项"),
        ("STEP 06", "发布与核验", "官网检查、审计确认"),
    ]
    x_start, box_width, gap, y, box_height = 65, 250, 34, 230, 190
    fills = ["#E8EEF5", "#E8EEF5", "#E8EEF5", "#FFF1DE", "#FFF1DE", "#E6F3EC"]
    outlines = ["#9AB7D1", "#9AB7D1", "#9AB7D1", "#E69B36", "#E69B36", "#56A178"]
    centers = []
    for index, ((stage, title, detail), fill, outline) in enumerate(zip(boxes, fills, outlines)):
        x = x_start + index * (box_width + gap)
        rect = (x, y, x + box_width, y + box_height)
        draw.rounded_rectangle(rect, radius=20, fill=fill, outline=outline, width=4)
        center_text(draw, (x + 16, y + 18, x + box_width - 16, y + 46), stage, stage_font, outline)
        center_text(draw, (x + 12, y + 58, x + box_width - 12, y + 112), title, label_font, "#0B2545")
        center_text(draw, (x + 14, y + 124, x + box_width - 14, y + 172), detail, detail_font, "#41566C")
        centers.append((x, x + box_width))
    for left, right in zip(centers, centers[1:]):
        start_x = left[1] + 8
        end_x = right[0] - 10
        mid_y = y + box_height // 2
        draw.line((start_x, mid_y, end_x, mid_y), fill="#7892A8", width=6)
        draw.polygon([(end_x, mid_y), (end_x - 18, mid_y - 11), (end_x - 18, mid_y + 11)], fill="#7892A8")

    draw.rounded_rectangle((94, 495, 822, 620), radius=18, fill="#FFF7EA", outline="#E8B55A", width=3)
    center_text(draw, (118, 514, 798, 551), "审核退回", label_font, "#7A5A00")
    center_text(draw, (118, 558, 798, 603), "根据审核意见回到“编辑草稿”，保存后重新预览和送审", note_font, "#5F4B21")
    draw.line((935, 420, 935, 480), fill="#E69B36", width=5)
    draw.polygon([(935, 480), (924, 460), (946, 460)], fill="#E69B36")

    draw.rounded_rectangle((978, 495, 1706, 620), radius=18, fill="#EFF7F2", outline="#77AF8B", width=3)
    center_text(draw, (1002, 514, 1682, 551), "发布路径", label_font, "#27613B")
    center_text(draw, (1002, 552, 1682, 607), "产品系列/型号/参数：产品统一发布工作台和快照；\n其他内容：内容发布队列", note_font, "#315A40", spacing=7)

    center_text(draw, (80, 672, 1720, 742), "任何内容错误均应通过“下线 -> 恢复为新草稿 -> 重新审核发布”处理，不直接改写已发布记录或数据库。", note_font, "#9B1C1C")
    canvas.save(FLOW_DIAGRAM, "PNG")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, before, after, color in ((1, 16, 18, 10, BLUE), (2, 13, 14, 7, BLUE), (3, 12, 10, 5, DARK_BLUE)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25
    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)


def configure_section(section):
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("瑞钧官网 CMS 运营标准化流程说明书")
    set_run_font(run, size=9, color=MUTED)
    footer = section.footer
    p = footer.paragraphs[0]
    add_page_number(p)


def set_core_properties(doc):
    props = doc.core_properties
    props.title = "瑞钧官网 CMS 运营标准化流程说明书"
    props.subject = "内容编辑、审核、发布和日常运维标准流程"
    props.author = "瑞钧官网运营团队"
    props.keywords = "CMS, Directus, 官网运营, SOP, 内容发布"


def make_document():
    create_flow_diagram()
    doc = Document()
    configure_styles(doc)
    set_core_properties(doc)
    section = doc.sections[0]
    configure_section(section)

    # Editorial cover: a restrained title page for an internal operating manual.
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("运营标准化流程说明书")
    set_run_font(r, size=12, color=GOLD, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("瑞钧官网 CMS")
    set_run_font(r, size=30, color=INK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("内容编辑、审核、发布与日常运维 SOP")
    set_run_font(r, size=15, color=DARK_BLUE)
    cover_rows = [
        ("版本", "v1.0"),
        ("生效日期", "2026-08-05"),
        ("适用对象", "内容编辑、技术/品牌审核、发布人员、系统管理员"),
        ("适用环境", "当前 Directus 本地演示环境；生产上线前需完成独立部署验收"),
    ]
    for label, value in cover_rows:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        label_run = p.add_run(f"{label}: ")
        set_run_font(label_run, size=10.5, color=INK, bold=True)
        value_run = p.add_run(value)
        set_run_font(value_run, size=10.5, color="27313D")
    doc.add_page_break()

    add_heading(doc, "1. 目的与适用范围")
    add_body(doc, "本说明书规定瑞钧官网 CMS 的日常内容维护流程，确保官网内容在编辑、审核、发布、下线与恢复过程中可追溯、可复核，并且不会将草稿、私有字段或未获授权素材暴露到官网。")
    add_callout(doc, "执行原则", "任何官网变更均遵循“编辑 -> 预览 -> 送审 -> 审核 -> 发布就绪检查 -> 发布 -> 发布后核验”的闭环。编辑人员只能维护草稿；审核、发布、下线和恢复由相应角色在专用工作台完成。")
    add_body(doc, "本 SOP 面向当前 Windows/SQLite 本地演示环境。其内容流程与权限边界可用于生产操作设计，但不等同于生产已具备 MySQL、对象存储、HTTPS、备份恢复、监控和正式域名等基础能力。")

    add_heading(doc, "2. 角色与职责")
    add_table(doc, ["角色", "主要职责", "可执行动作", "不得执行"], [
        ("内容编辑", "编写和维护业务内容", "新建/保存草稿、发起预览、提交审核、补充审核意见", "直接发布、绕过受控媒体、修改已发布记录"),
        ("技术审核", "确认参数、服务资料、FAQ 安全边界与技术主张", "审核、退回并写明意见", "代替发布人员发布内容"),
        ("品牌/业务审核", "确认品牌表达、客户授权、版权与宣传表述", "审核、退回并写明意见", "删除审计记录或改写技术结论"),
        ("发布人员", "执行发布、下线、归档与发布后核验", "查看就绪检查、正式发布、下线、归档", "在状态迁移中修改正文、参数或媒体"),
        ("系统管理员", "账号、权限、运行状态和维护窗口管理", "配置权限、运行自检、扩展更新与重启", "以管理员身份代替业务审核结论"),
    ], [1600, 2600, 2700, 2460])

    add_heading(doc, "3. 日常进入内容编辑工作台", 1)
    add_heading(doc, "3.1 启动与访问", 2)
    add_number(doc, "在具备授权的本机或内网环境确认 CMS 服务已启动。日常本地启动由系统管理员在 cms 目录执行运行脚本；普通运营人员不修改 .env.local、数据库文件或服务端密钥。")
    add_code(doc, r"Set-Location D:\CURSORpj\gaunwang\cms")
    add_code(doc, r"powershell -ExecutionPolicy Bypass -File .\scripts\run-local-directus.ps1")
    add_number(doc, "使用浏览器访问 CMS 管理端。当前本机入口为 http://127.0.0.1:8055/admin/；获授权的内网访问可使用项目配置的内网地址。")
    add_number(doc, "使用分配的个人账号登录。账号、密码、令牌和 .env.local 中的任何密钥不得写入工单、聊天记录、截图、共享文档或源码。")
    add_number(doc, "打开内容编辑工作台：可从后台侧栏选择“内容编辑工作台”，或直接访问 /admin/ruijun-content-editor-workbench。直接地址用于侧栏未刷新、浏览器书签或培训演示，不应共享含个人会话的信息。")
    add_callout(doc, "推荐入口", "http://127.0.0.1:8055/admin/ruijun-content-editor-workbench。若在内网使用，请将主机替换为系统管理员发布的授权地址，端口保持与运行配置一致。")

    add_figure(doc, WORKBENCH_HOME, "图 1  内容编辑工作台首页：从任务式入口开始日常内容更新。", "瑞钧官网 CMS 内容编辑工作台首页，展示编辑、预览、送审三步和五项常用任务入口。")
    add_heading(doc, "3.2 工作台功能分区", 2)
    add_table(doc, ["区域", "日常用途", "操作提示"], [
        ("常用内容", "页面文案、产品系列、企业资料、服务支持、文章与案例", "首次进入从任务式起始页选择内容类型；只维护业务字段。"),
        ("更多工具", "媒体资产、FAQ 知识、型号与参数、提交审核、审核队列、全站设置", "仅在当前任务需要时打开；技术字段保留在高级入口。"),
        ("顶部操作栏", "重新加载、查看操作说明、预览当前草稿、在官网中预览草稿", "存在“有未保存修改”提示时，先保存或明确放弃。"),
    ], [1800, 3600, 3960])

    add_figure(doc, WORKBENCH_TOOLS, "图 2  展开“更多工具”后，可进入媒体、送审、审核和全站设置等辅助工作区。", "内容编辑工作台中展开更多工具菜单的界面，显示媒体资产、常见问题知识、型号与参数、提交审核、审核队列和全站设置。")
    add_heading(doc, "4. 标准内容变更流程")
    add_callout(doc, "流程总览", "变更申请与资料准备 -> 草稿编辑 -> 受控媒体处理 -> 保存 -> 双重预览 -> 提交审核 -> 审核处理 -> 发布就绪检查 -> 正式发布 -> 官网核验与审计归档。")
    add_figure(doc, FLOW_DIAGRAM, "图 3  官网内容变更闭环：审核退回后回到草稿；产品与其他内容使用不同的发布路径。", "官网内容变更流程图，从准备依据、编辑草稿、双重预览、提交审核到审核就绪、发布和官网核验，并标出退回和产品发布分支。", width=6.35)
    add_heading(doc, "4.1 变更准备", 2)
    add_number(doc, "明确变更目标、内容归属、计划上线时间、适用语言和责任人。涉及产品参数、资质、服务网点、售后入口、客户案例或宣传主张时，先取得对应业务或技术依据。")
    add_number(doc, "准备可公开使用的文字和素材。图片、视频、PDF 必须明确使用范围、替代文本、版权状态和来源/授权说明；不确定授权状态的素材不得进入官网内容。")
    add_number(doc, "判断是否属于产品统一发布。产品系列、产品型号和产品参数必须进入“产品统一发布工作台”，不能通过通用内容发布队列单独发布。")

    add_heading(doc, "4.2 编辑草稿", 2)
    add_number(doc, "在“常用内容”选择页面文案、产品系列、企业资料、服务支持或文章与案例。仅编辑面向访客的业务字段；路由标识、来源记录和内部段落标识位于“高级设置（一般无需修改）”，除非有明确授权，否则不修改。")
    add_number(doc, "填写必填字段，并保持标题、摘要、正文、分类、SEO、来源和关联关系一致。文章可使用正文或视频地址，但发布前仍需要摘要、封面、SEO 与来源。")
    add_number(doc, "点击保存。保存只更新草稿，不会对外发布；状态为审核中、已排期或已发布的记录不可在编辑台直接改写，应先按版本流程恢复为草稿。")
    add_callout(doc, "保存边界", "工作台保存不等于发布。任何未保存修改在切换内容、重新加载或关闭后台前都会触发确认；不要依赖浏览器自动恢复。", color=GOLD)

    add_figure(doc, PAGE_FORM, "图 4  页面文案编辑示例：左侧选择草稿，右侧填写访客可见的标题、段落和搜索信息。", "页面文案编辑表单，左侧为草稿列表，右侧为页面标题、段落、搜索与分享信息和保存草稿操作。")
    add_heading(doc, "4.3 处理受控媒体", 2)
    add_number(doc, "需要新增图片、视频或 PDF 时，在“更多工具 -> 媒体资产”上传并登记候选。填写文件、使用范围、替代文本、版权状态和授权/来源说明。")
    add_number(doc, "候选素材先保持私有草稿。对应技术或品牌审核人员完成审核后，由发布人员发布素材；素材达到已发布状态前，不会出现在内容表单的受控选择器中。")
    add_number(doc, "返回内容草稿，按用途选择已发布媒体资产：产品使用 product，文章使用 article，服务资料使用 service，品牌 Logo 使用 brand，资质证书使用 qualification，制造证据使用 manufacturing。")
    add_callout(doc, "禁止事项", "不得直接填写文件 ID、使用本地路径或绕过媒体审核。文件上传成功也不代表已获官网使用资格。", color=RED)

    add_heading(doc, "4.4 双重预览与送审", 2)
    add_number(doc, "先点击“预览当前草稿”，检查未保存表单中的文字、字段组织和基本展示。该预览不保存、不发布，也不生成公开链接。")
    add_number(doc, "保存草稿后，点击“在官网中预览草稿”，在官网真实版式中复核标题层级、图片比例、手机端阅读、链接和 CTA。预览由短时效的一次性授权流程承载，浏览器地址不应包含草稿正文或凭据。")
    add_number(doc, "预览无误后进入“提交审核”，填写变更说明、依据来源、影响范围和期望上线时间。送审后不要以本地保存替代审核结论。")
    add_number(doc, "审核人员在“审核队列”核对业务、技术、版权、SEO、素材和风险边界。通过后内容进入待发布/已排期状态；退回时必须根据审核意见恢复草稿、修订、重新预览并再次送审。")

    add_heading(doc, "4.5 发布前检查与正式发布", 2)
    add_number(doc, "发布人员进入“发布前就绪检查”，筛选待发布内容，处理所有阻断项。就绪检查只显示内容标签、阻断项和原生详情入口，不会授予额外编辑权限。")
    add_number(doc, "非产品内容进入“内容发布队列”，在“待发布”中选择已排期记录，填写发布说明后执行正式发布。状态迁移只提交目标状态和说明，不能趁机修改正文、参数或媒体。")
    add_number(doc, "产品系列、型号和参数进入“产品统一发布工作台”。发布前确认系列、型号和参数关联完整；系统会生成不可变产品快照，官网列表、详情与对比均读取同一版本。")
    add_number(doc, "发布后确认记录进入已发布状态，并在官网核验公开页面。若发现问题，优先使用下线或恢复为草稿的受控流程，不直接修改已发布内容或数据库。")

    add_heading(doc, "5. 发布后核验与变更关闭")
    add_table(doc, ["核验项", "责任人", "完成标准"], [
        ("公开内容", "发布人员", "官网页面、列表、详情、搜索和相关入口展示同一已发布版本。"),
        ("产品内容", "发布人员 + 技术审核", "产品列表、详情、对比均来自同一产品快照；参数和型号关系一致。"),
        ("链接与 CTA", "内容编辑 + 发布人员", "站内路径正确；外部服务入口可用且保留人工联系兜底。"),
        ("媒体与可读性", "内容编辑", "图片/视频/PDF 加载正确，替代文本与授权信息完整，移动端未出现明显截断。"),
        ("审计与回退", "发布人员", "发布说明、审核人、时间和版本记录存在；已知回退路径已确认。"),
    ], [2300, 2400, 4660])
    add_body(doc, "变更关闭前，在变更记录中注明：发布内容、公开链接、上线时间、核验人、核验结果、异常与后续动作。无异常时关闭变更；存在问题时转入异常处置并保留审计链路。")

    add_heading(doc, "6. 下线、归档与版本恢复")
    add_number(doc, "需要撤回已公开的非产品内容时，发布人员在“内容发布队列 -> 已发布”执行下线。下线后内容不应继续通过公开 API、站内搜索或历史 URL 对外可见。")
    add_number(doc, "确认无需保留为可恢复业务草稿时，在“已下线”执行归档。归档保留审计轨迹，不等同于删除数据库记录。")
    add_number(doc, "需要回退内容时，进入“内容版本”工作台，选择历史版本并填写恢复原因。系统会创建新的未发布草稿并记录操作者、时间、来源版本和差异，恢复后仍需重新审核与发布。")
    add_number(doc, "需要回退产品时，在“产品统一发布工作台”选择历史快照并填写恢复原因。系统将复制为新的递增版本，而不是重新激活旧快照；恢复后复核官网缓存刷新状态。")

    add_heading(doc, "7. 日常运维节奏")
    add_table(doc, ["频次", "动作", "输出/记录"], [
        ("每次发布", "执行预览、就绪检查、发布后官网核验和版本确认。", "变更记录、审核意见、发布说明、核验结果。"),
        ("每日", "查看运营总览中的草稿、待审核、已发布和待处理数量；跟进临近排期与退回项。", "待办清单及责任人。"),
        ("每周", "核对媒体候选、待审核内容、下线记录、服务资料版本和外链健康状态。", "周度内容质量清单。"),
        ("每月", "复核账号权限、管理员数量、内容版本恢复记录、官网缓存异常和备份恢复演练状态。", "月度运维记录与风险项。"),
        ("维护窗口", "系统管理员执行只读运行自检，必要时构建扩展、重启服务并回归关键工作台。", "维护变更单、运行自检结果、回归结论。"),
    ], [1400, 4380, 3580])

    add_heading(doc, "8. 常用工作台与路径")
    add_table(doc, ["工作台", "主要用途", "后台路径"], [
        ("内容编辑工作台", "日常草稿编辑、预览、送审入口", "/admin/ruijun-content-editor-workbench"),
        ("运营总览", "只读查看内容状态与待办，并跳转至相应工作台", "/admin/ruijun-operations-overview-workbench"),
        ("发布前就绪检查", "发现发布阻断项并跳转到原生内容详情", "/admin/ruijun-publication-readiness-workbench"),
        ("内容发布队列", "非产品内容的发布、下线与归档", "/admin/ruijun-content-publication-queue-workbench"),
        ("产品统一发布", "产品快照发布与历史快照恢复", "/admin/ruijun-product-release-workbench"),
        ("内容版本", "查看差异和将历史版本恢复为新草稿", "/admin/ruijun-content-version-manager"),
        ("FAQ 审核", "查看 FAQ 技术审核摘要和风险待办", "/admin/ruijun-knowledge-review-workbench"),
    ], [1950, 3800, 3610])

    add_heading(doc, "9. 异常处置")
    add_table(doc, ["情形", "立即动作", "后续处置", "禁止动作"], [
        ("CMS 无法访问", "停止编辑；记录时间、访问地址和错误提示。", "系统管理员执行运行自检，确认服务健康、扩展构建和网络范围。", "反复重启、修改数据库或公开凭据。"),
        ("发现发布内容错误", "立即由发布人员下线，保留发布说明和证据。", "从版本恢复为新草稿，修订后重新审核、发布和核验。", "直接修改已发布记录或删除审计。"),
        ("媒体无权使用/失效", "停止引用并下线受影响内容或素材。", "补充授权材料，重新走媒体候选、审核、发布流程。", "用文件路径或临时外链替代受控素材。"),
        ("发布就绪检查阻断", "不发布；记录阻断项。", "回到原生内容详情补齐字段、关系、SEO、来源或审核材料。", "忽略阻断项或以管理员权限绕过。"),
        ("官网仍显示旧内容", "确认 CMS 已发布版本与官网页面版本。", "检查官网缓存失效状态；由管理员按既定维护流程处理并复核。", "通过修改草稿或重复发布碰运气。"),
    ], [1800, 2500, 3300, 1760])

    add_heading(doc, "10. 安全与审计要求")
    add_bullet(doc, "使用个人账号完成操作，不共享账号，不借用发布人员或管理员会话。")
    add_bullet(doc, "账号密码、管理员令牌、服务令牌、Webhook 密钥和 .env.local 仅可保存在授权的密钥管理或本地受控环境，禁止写入本说明书、截图、前端代码或提交记录。")
    add_bullet(doc, "不直接编辑 SQLite/MySQL 数据库，不通过 API 工具绕过状态机，不以本地开发环境代替生产发布审批。")
    add_bullet(doc, "每次审核、发布、下线、归档和恢复必须填写说明，确保后续可追溯变更原因、操作者与时间。")
    add_bullet(doc, "生产上线前，必须完成 HTTPS、最小权限、对象存储、备份恢复、监控告警、会话与预览令牌生产化验收。")

    add_heading(doc, "附录 A. 发布前检查清单")
    checklist = [
        "变更目标、责任人、来源依据和上线窗口已明确。",
        "内容已保存为草稿，且必填字段、SEO、关联关系、语言和公开范围完整。",
        "图片、视频和 PDF 均为已发布的受控媒体资产，使用范围、替代文本、版权状态与授权说明无误。",
        "已分别完成“预览当前草稿”和“在官网中预览草稿”的复核。",
        "审核意见已处理，技术、品牌/业务与客户授权边界已确认。",
        "发布前就绪检查无阻断项；产品内容已在产品统一发布工作台核对快照。",
        "发布说明、发布后核验人和回退路径已准备。",
    ]
    for item in checklist:
        add_bullet(doc, item)

    add_heading(doc, "附录 B. 系统管理员维护检查")
    add_body(doc, "仅在维护窗口执行。运行自检不写数据库、不启动或停止服务，可用于确认本地 Directus 健康、原生依赖和扩展构建状态。")
    add_code(doc, r"Set-Location D:\CURSORpj\gaunwang\cms")
    add_code(doc, r"npm run runtime:verify")
    add_body(doc, "如更新了内容编辑工作台扩展，先构建扩展，再在维护窗口重启 Directus。重启后回归内容编辑、发布就绪检查和发布队列等关键入口。")
    add_code(doc, r"Set-Location .\extensions\content-editor-workbench")
    add_code(doc, r"npm run build")
    add_code(doc, r"Set-Location ..\..")
    add_code(doc, r"npm run demo:restart")
    add_callout(doc, "维护边界", "demo:seed 仅用于本地演示临时草稿；演示结束应运行 demo:cleanup。不得将演示数据当作官网正式业务内容，也不得将本地 SQLite 数据用于生产。", color=GOLD)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    make_document()
